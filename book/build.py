import glob,re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
d=Document()
st=d.styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(12)
for f in sorted(glob.glob('part*.txt')):
    for l in open(f,encoding='utf-8').read().splitlines():
        if not l.strip(): continue
        if l.startswith('#T '): p=d.add_paragraph(); p.alignment=A.CENTER; r=p.add_run(l[3:]); r.font.size=Pt(20)
        elif l.startswith('#C '): p=d.add_paragraph(); p.alignment=A.CENTER; r=p.add_run(l[3:]); r.font.size=Pt(16)
        elif l.startswith('#H '): p=d.add_paragraph(); p.add_run(l[3:]).bold=True
        else: p=d.add_paragraph(l); p.alignment=A.JUSTIFY
d.save('Магия_шнура_и_зигзага.docx')
